"""KOSIF Sheets — build Excel workbooks and SQLite databases as files (no GUI driving).

* `kosif_sheet_build`  one call → a finished .xlsx: sheets with a real Excel Table (filter buttons, banded style),
  rows sorted the way you asked, number formats, auto column widths, frozen header, optional totals row with
  live SUM formulas, right-to-left sheets for Arabic, and an optional bar/line/pie chart.
* `kosif_db`           SQLite: create tables from a typed schema, insert rows (parameterised), run read-only
  SELECT queries (the connection itself is opened read-only), describe the database, export a query to .xlsx/.csv.

Files are written to Documents\\KOSIF (or inside a folder you enabled for KOSIF). `open: true` opens the finished
file in its default program (Excel) so you can see it.
"""
from __future__ import annotations

import csv
import datetime as dt
import os
import re
import sqlite3
from contextlib import contextmanager
from pathlib import Path

from .paths import SandboxError, allowed_roots, _inside
from .registry import A, B, I, O, S, ToolError, obj, tool

TYPES = {"integer": "INTEGER", "int": "INTEGER", "real": "REAL", "float": "REAL", "number": "REAL",
         "numeric": "NUMERIC", "text": "TEXT", "string": "TEXT", "date": "TEXT", "datetime": "TEXT",
         "bool": "INTEGER", "boolean": "INTEGER"}
FORMATS = {"money": '#,##0.00', "integer": '#,##0', "number": '#,##0.00', "percent": '0.0%', "date": 'yyyy-mm-dd',
           "datetime": 'yyyy-mm-dd hh:mm', "text": '@'}
IDENT = re.compile(r"^[A-Za-z_؀-ۿ][\w؀-ۿ ]{0,62}$")


# ── where files may go ───────────────────────────────────────────────────────
def out_dir() -> Path:
    d = Path(os.environ.get("KOSIF_OFFICE_DIR") or (Path.home() / "Documents" / "KOSIF")).expanduser()
    d.mkdir(parents=True, exist_ok=True)
    return d.resolve()


def _safe_name(name: str, ext: str) -> str:
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", Path(str(name)).stem).strip(" .") or "kosif"
    return stem[:80] + ext


def target_path(path: str | None, default_name: str, ext: str, *, must_exist: bool = False) -> Path:
    """A path inside Documents\\KOSIF or a KOSIF-enabled folder; never anywhere else."""
    roots = [out_dir(), *allowed_roots()]
    if not path:
        p = out_dir() / _safe_name(default_name, ext)
    else:
        p = Path(path).expanduser()
        if not p.is_absolute():
            p = out_dir() / p
        if p.suffix.lower() != ext:
            p = p.with_suffix(ext)
    p = p.resolve()
    if not any(_inside(p, r) for r in roots):
        raise SandboxError(f"{p} is outside Documents\\KOSIF and the folders enabled for KOSIF")
    if must_exist and not p.exists():
        raise ToolError(f"not found: {p}")
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _open(p: Path) -> bool:
    try:
        os.startfile(str(p))  # noqa: S606 - opens the file the user asked us to build, in its default app
        return True
    except (AttributeError, OSError):
        return False


# ── Excel ────────────────────────────────────────────────────────────────────
def _coerce(v):
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"-?\d+", s) and len(s) < 16 and not (len(s) > 1 and s.startswith("0")):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
            try:
                return dt.date.fromisoformat(s)
            except ValueError:
                return v
    return v


def _rows_as_lists(headers: list[str], rows: list) -> list[list]:
    out = []
    for r in rows or []:
        if isinstance(r, dict):
            out.append([_coerce(r.get(h)) for h in headers])
        elif isinstance(r, (list, tuple)):
            out.append([_coerce(x) for x in list(r)[:len(headers)]] + [None] * max(0, len(headers) - len(r)))
        else:
            raise ToolError("each row must be an object {header: value} or a list in header order")
    return out


def _sortkey(v):
    """Numbers, then dates, then text (case-insensitive); mixed columns never raise."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return (0, v, "")
    if isinstance(v, (dt.date, dt.datetime)):
        return (1, 0, v.isoformat())
    return (2, 0, str(v).lower())


def _sort(headers: list[str], rows: list[list], sort) -> list[list]:
    keys = sort if isinstance(sort, list) else [sort] if sort else []
    for k in reversed(keys):  # stable sorts applied last key first = multi-key sort
        if isinstance(k, str):
            k = {"column": k}
        col = k.get("column")
        if col not in headers:
            raise ToolError(f"sort column {col!r} is not a header; headers: {headers}")
        i = headers.index(col)
        present = [r for r in rows if r[i] not in (None, "")]
        empty = [r for r in rows if r[i] in (None, "")]
        present.sort(key=lambda r: _sortkey(r[i]), reverse=bool(k.get("desc")))
        rows = present + empty  # empty cells always last
    return rows


def _table_name(name: str, used: set) -> str:
    """Excel table names must be ASCII-ish identifiers: keep Latin names, number the rest (Table1, Table2…)."""
    base = re.sub(r"[^A-Za-z0-9_]", "_", name).strip("_")
    if not re.search(r"[A-Za-z]", base):
        base = "Table"
    base = ("T_" + base) if not base[0].isalpha() else base
    if base == "Table":
        i = 1
        while f"table{i}" in used:
            i += 1
        used.add(f"table{i}")
        return f"Table{i}"
    n, i = base[:240], 2
    while n.lower() in used:
        n, i = f"{base[:236]}_{i}", i + 1
    used.add(n.lower())
    return n


def build_workbook(sheets: list[dict], *, path: str | None = None, title: str = "kosif", open_after: bool = False,
                   overwrite: bool = False) -> dict:
    try:
        from openpyxl import Workbook  # noqa: PLC0415
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference  # noqa: PLC0415
        from openpyxl.styles import Alignment, Font  # noqa: PLC0415
        from openpyxl.utils import get_column_letter  # noqa: PLC0415
        from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: PLC0415
    except ImportError:
        raise ToolError("Excel files need the `openpyxl` package (pip install openpyxl)") from None
    if not sheets:
        raise ToolError("give at least one sheet: [{name, headers, rows}]")
    p = target_path(path, title, ".xlsx")
    if p.exists() and not overwrite:
        raise ToolError(f"{p.name} already exists: pass overwrite=true or another path")
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    report = []
    for si, sh in enumerate(sheets):
        name = re.sub(r"[\[\]:*?/\\]", "_", str(sh.get("name") or f"Sheet{si + 1}"))[:31]
        headers = [str(h) for h in (sh.get("headers") or [])]
        if not headers and sh.get("rows") and isinstance(sh["rows"][0], dict):
            headers = list(sh["rows"][0].keys())
        if not headers:
            raise ToolError(f"sheet {name!r}: headers are required")
        if len(set(headers)) != len(headers):
            raise ToolError(f"sheet {name!r}: duplicate headers")
        rows = _sort(headers, _rows_as_lists(headers, sh.get("rows") or []), sh.get("sort"))
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = bool(sh.get("rtl", any(re.search("[؀-ۿ]", h) for h in headers)))
        ws.append(headers)
        for r in rows:
            ws.append(r)
        n = len(rows)
        last_col = get_column_letter(len(headers))
        fmts = sh.get("formats") or {}
        for ci, h in enumerate(headers, 1):
            f = fmts.get(h)
            if f:
                code = FORMATS.get(f, f)
                for (cell,) in ws.iter_rows(min_row=2, max_row=n + 1, min_col=ci, max_col=ci):
                    cell.number_format = code
            elif n and all(isinstance(r[ci - 1], dt.date) for r in rows if r[ci - 1] is not None):
                for (cell,) in ws.iter_rows(min_row=2, max_row=n + 1, min_col=ci, max_col=ci):
                    cell.number_format = FORMATS["date"]
        if n:
            tname = _table_name(sh.get("table") or name, used)
            t = Table(displayName=tname, ref=f"A1:{last_col}{n + 1}")
            t.tableStyleInfo = TableStyleInfo(name=sh.get("style") or "TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        else:
            tname = None
        totals = sh.get("totals")
        if totals and n:
            cols = totals if isinstance(totals, list) else [h for i, h in enumerate(headers)
                                                             if all(isinstance(r[i], (int, float)) and not isinstance(r[i], bool)
                                                                    for r in rows if r[i] is not None)
                                                             and any(r[i] is not None for r in rows)]
            tr = n + 2
            ws.cell(row=tr, column=1, value=sh.get("totals_label") or ("الإجمالي" if ws.sheet_view.rightToLeft else "Total"))
            ws.cell(row=tr, column=1).font = Font(bold=True)
            for h in cols:
                if h not in headers:
                    raise ToolError(f"totals column {h!r} is not a header")
                ci = headers.index(h) + 1
                if ci == 1:
                    continue
                L = get_column_letter(ci)
                c = ws.cell(row=tr, column=ci, value=f"=SUM({L}2:{L}{n + 1})")
                c.font = Font(bold=True)
                if fmts.get(h):
                    c.number_format = FORMATS.get(fmts[h], fmts[h])
        for ci, h in enumerate(headers, 1):  # auto width
            vals = [h] + [r[ci - 1] for r in rows[:500]]
            width = max(len(str(v)) if v is not None else 0 for v in vals)
            ws.column_dimensions[get_column_letter(ci)].width = min(max(10, width + 3), 60)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        ch = sh.get("chart")
        if ch and n:
            kind = (ch.get("type") or "bar").lower()
            label = ch.get("labels") or headers[0]
            values = ch.get("values") or [h for h in headers if h != label][:1]
            if label not in headers or any(v not in headers for v in values):
                raise ToolError("chart labels/values must be header names")
            chart = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart}.get(kind, BarChart)()
            chart.title = ch.get("title") or ", ".join(values)
            if kind == "bar":
                chart.type = "bar"
            li = headers.index(label) + 1
            for v in values:
                vi = headers.index(v) + 1
                chart.add_data(Reference(ws, min_col=vi, min_row=1, max_row=n + 1), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=li, min_row=2, max_row=n + 1))
            chart.width, chart.height = 18, 9
            ws.add_chart(chart, f"{get_column_letter(len(headers) + 2)}2")
        report.append({"sheet": name, "rows": n, "columns": len(headers), "table": tname,
                       "sorted_by": sh.get("sort"), "totals": bool(totals), "chart": bool(ch), "rtl": ws.sheet_view.rightToLeft})
    wb.save(p)
    out = {"file": str(p), "bytes": p.stat().st_size, "sheets": report}
    if open_after:
        out["opened"] = _open(p)
    return out


# ── SQLite ───────────────────────────────────────────────────────────────────
def _sql(template: str, **names) -> str:
    """The ONLY place SQL text is assembled: a fixed template whose {slots} take identifiers that passed `_q`
    (strict name pattern + double-quote escaping). A list fills a slot with a comma-separated name list.
    Values never go through here — they are always bound as `?` parameters."""
    return template.format(**{k: ", ".join(_q(x) for x in v) if isinstance(v, (list, tuple)) else _q(v)
                              for k, v in names.items()})


def _q(ident: str) -> str:
    s = str(ident).strip()
    if not IDENT.match(s):
        raise ToolError(f"invalid name {ident!r}: letters, digits, spaces and _ only (max 63)")
    return '"' + s.replace('"', '""') + '"'


@contextmanager
def _connect(p: Path, *, readonly: bool):
    """Commit on success, roll back on error, and always close the file handle."""
    c = sqlite3.connect(f"file:{p.as_posix()}?mode=ro", uri=True) if readonly else sqlite3.connect(p)
    try:
        if not readonly:
            c.execute("PRAGMA foreign_keys=ON")
        yield c
        if not readonly:
            c.commit()
    except Exception:
        c.rollback()
        raise
    finally:
        c.close()


READ_SQL = re.compile(r"^\s*(select|with)\b", re.I)


def db(op: str, a: dict) -> dict:
    name = a.get("db") or "kosif"
    p = target_path(name if str(name).lower().endswith(".db") or os.sep in str(name) or "/" in str(name) else None,
                    str(name), ".db", must_exist=op in ("query", "describe", "export"))
    if op == "create":
        table, cols = a.get("table"), a.get("columns") or {}
        if not table or not cols:
            raise ToolError("create needs `table` and `columns` {name: type}")
        pk = a.get("primary_key")
        parts = []
        for cname, ctype in cols.items():
            t = TYPES.get(str(ctype).lower())
            if not t:
                raise ToolError(f"column {cname!r}: type {ctype!r} not in {sorted(set(TYPES))}")
            parts.append(f"{_q(cname)} {t}" + (" PRIMARY KEY" if pk == cname else "") +
                         (" NOT NULL" if cname in (a.get("required") or []) else "") +
                         (" UNIQUE" if cname in (a.get("unique") or []) else ""))
        with _connect(p, readonly=False) as c:
            if a.get("replace"):
                c.execute(_sql("DROP TABLE IF EXISTS {t}", t=table))
            create = _sql("CREATE TABLE IF NOT EXISTS {t} ", t=table) + "(" + ", ".join(parts) + ")"
            c.execute(create)  # parts = quoted names + types/constraints from fixed allowlists
            for col in a.get("index") or []:
                c.execute(_sql("CREATE INDEX IF NOT EXISTS {i} ON {t} ({c})",
                               i="idx_" + str(table) + "_" + str(col), t=table, c=col))
        return {"db": str(p), "table": table, "columns": cols, "created": True}
    if op == "insert":
        table, rows = a.get("table"), a.get("rows") or []
        if not table or not rows:
            raise ToolError("insert needs `table` and `rows`")
        with _connect(p, readonly=False) as c:
            cols = [r[1] for r in c.execute(_sql("PRAGMA table_info({t})", t=table))]
            if not cols:
                raise ToolError(f"table {table!r} does not exist: create it first")
            unknown = sorted({k for r in rows if isinstance(r, dict) for k in r} - set(cols))
            if unknown:
                raise ToolError(f"columns {unknown} are not in table {table!r} ({cols})")
            n = 0
            for r in rows:  # objects map by column name, lists by column order; both may be mixed
                if isinstance(r, dict):
                    use, vals = list(r.keys()), list(r.values())
                elif isinstance(r, (list, tuple)):
                    if len(r) > len(cols):
                        raise ToolError(f"row {r!r} has more values than table {table!r} has columns")
                    use, vals = cols[:len(r)], list(r)
                else:
                    raise ToolError("each row must be an object or a list")
                stmt = _sql("INSERT INTO {t} ({cols}) ", t=table, cols=use) + "VALUES (" + ", ".join("?" * len(use)) + ")"
                c.execute(stmt, vals)
                n += 1
            total = c.execute(_sql("SELECT COUNT(*) FROM {t}", t=table)).fetchone()[0]
        return {"db": str(p), "table": table, "inserted": n, "rows_now": total}
    if op in ("query", "export"):
        sql = str(a.get("sql") or "").strip().rstrip(";")
        if not READ_SQL.match(sql) or ";" in sql:
            raise ToolError("query/export take ONE read-only SELECT (or WITH … SELECT) statement")
        limit = int(a.get("limit", 1000))
        with _connect(p, readonly=True) as c:
            cur = c.execute(sql, a.get("params") or [])
            headers = [d[0] for d in cur.description or []]
            rows = cur.fetchmany(limit if op == "query" else 1_000_000)
        if op == "query":
            return {"db": str(p), "columns": headers, "rows": [list(r) for r in rows], "count": len(rows),
                    "truncated": len(rows) >= limit}
        fmt = (a.get("format") or "xlsx").lower()
        if fmt == "csv":
            out = target_path(a.get("path"), a.get("name") or Path(p).stem, ".csv")
            with open(out, "w", newline="", encoding="utf-8-sig") as fh:
                w = csv.writer(fh)
                w.writerow(headers)
                w.writerows(rows)
            res = {"file": str(out), "rows": len(rows)}
            if a.get("open"):
                res["opened"] = _open(out)
            return res
        return build_workbook([{"name": a.get("sheet") or "Data", "headers": headers, "rows": [list(r) for r in rows],
                                "formats": a.get("formats"), "totals": a.get("totals"), "chart": a.get("chart")}],
                              path=a.get("path"), title=a.get("name") or Path(p).stem, open_after=bool(a.get("open")),
                              overwrite=bool(a.get("overwrite")))
    if op == "describe":
        with _connect(p, readonly=True) as c:
            tables = [r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
            out = {}
            for t in tables:
                info = c.execute(_sql("PRAGMA table_info({t})", t=t)).fetchall()
                out[t] = {"columns": [{"name": r[1], "type": r[2], "pk": bool(r[5])} for r in info],
                          "rows": c.execute(_sql("SELECT COUNT(*) FROM {t}", t=t)).fetchone()[0]}
        return {"db": str(p), "tables": out}
    raise ToolError("op must be create | insert | query | describe | export")


# ── MCP tools ────────────────────────────────────────────────────────────────
@tool("kosif_sheet_build",
      "KOSIF Sheets — build a finished Excel workbook (.xlsx) in one call: per sheet a real Excel Table (filters, "
      "banded style), rows sorted as asked, number formats (money/percent/date/integer or an Excel code), auto "
      "widths, frozen header, optional totals row with live SUM formulas, right-to-left for Arabic, optional "
      "bar/line/pie chart. Saved to Documents\\KOSIF (or an enabled folder); open=true opens it in Excel.",
      obj({"sheets": A(O("{name, headers:[...], rows:[{header:value}|[...]], sort:[{column, desc}], "
                         "formats:{header:'money'|'percent'|'date'|'integer'|code}, totals:true|[headers], "
                         "chart:{type:'bar'|'column'|'line'|'pie', labels:header, values:[headers], title}, "
                         "rtl:bool, style:'TableStyleMedium2', table:'Name'}"), "Sheets in order."),
           "path": S("File name or path (.xlsx). Default: <title>.xlsx in Documents\\KOSIF."),
           "title": S("Default file name."), "open": B("Open the file in Excel when done."),
           "overwrite": B("Replace an existing file.")}, ["sheets"]),
      title="Build Excel workbook", read_only=False, idempotent=False)
def _t_sheet(a: dict) -> dict:
    return build_workbook(a["sheets"], path=a.get("path"), title=a.get("title") or "kosif",
                          open_after=bool(a.get("open")), overwrite=bool(a.get("overwrite")))


@tool("kosif_db",
      "KOSIF Sheets — a SQLite database file in Documents\\KOSIF: op=create (table + typed columns, primary key, "
      "unique/required, indexes), insert (rows as objects or lists, parameterised), query (ONE read-only SELECT, "
      "connection opened read-only), describe (tables, columns, row counts), export (a SELECT → formatted .xlsx "
      "or .csv, optional open).",
      obj({"op": S("create | insert | query | describe | export", enum=["create", "insert", "query", "describe", "export"]),
           "db": S("Database name or path (default 'kosif' → kosif.db)."), "table": S(),
           "columns": O("{column: integer|real|numeric|text|date|datetime|boolean}"),
           "primary_key": S(), "required": A(S()), "unique": A(S()), "index": A(S()),
           "replace": B("create: drop the table first."), "rows": A({}, "insert rows"),
           "sql": S("query/export: one SELECT statement"), "params": A({}, "query parameters for ? placeholders"),
           "limit": I("query: max rows (default 1000)", minimum=1, maximum=100000),
           "format": S("export: xlsx (default) | csv", enum=["xlsx", "csv"]), "path": S("export file"),
           "name": S("export file name"), "sheet": S(), "formats": O(), "totals": {"description": "true or [headers]"},
           "chart": O(), "open": B("export: open the file"), "overwrite": B()}, ["op"]),
      title="SQLite database", read_only=False, idempotent=False)
def _t_db(a: dict) -> dict:
    return db(str(a["op"]), a)


__all__ = ["build_workbook", "db", "out_dir", "target_path"]
